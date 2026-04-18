# app/main.py

from __future__ import annotations

from typing import List
import os
import tempfile
from io import BytesIO

import pandas as pd
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.encoders import jsonable_encoder


from .core import parse_one_pdf, LABELS_FLAT, build_sales_dataframe, filename_sort_key
from m1_pipeline.parse_check import build_parse_check
from m1_pipeline.device_id import attach_device_ids
from m1_pipeline.device_status import build_device_status, build_device_status_from_existing_status
from m1_pipeline.issuance_status import build_issuance_status
from m1_pipeline.cost_redemption import attach_redemption_cost
from m1_pipeline.transfer_vintage import attach_transfer_vintage
from m1_pipeline.device_wise_sales import fill_device_wise_sales
from m1_pipeline.cogs import compute_cogs
from .finance_agent import answer_finance


app = FastAPI(
    title="Project Max Certificates API",
    version="0.2.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],          # you can restrict later
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/health")
def health():
    return {"status": "ok"}

@app.get("/")
def root():
    # Used by Render's health check
    return {"status": "ok", "service": "project-max-backend"}


# in app/main.py

from fastapi import Form

@app.post("/m1/issuance-from-is/xlsx")
async def issuance_from_is_xlsx(
    is_file: UploadFile = File(..., description="I/S master Excel"),
    source_mode: str = Form("external"),  # "external" or "sheet"
    # optional external file (used only if source_mode == "external")
    issuance_file: UploadFile | None = File(None, description="Issuance Status Excel"),
):
    """
    Build Issuance Status sheet from:
      - either a separate Issuance file (external)
      - or the 'Issuance Status' sheet inside the I/S Excel
    """
    import pandas as pd
    from io import BytesIO
    from m1_pipeline.issuance_status import build_issuance_status

    # Read I/S Excel bytes once
    is_bytes = await is_file.read()

    if source_mode == "sheet":
        # 👉 read sheet 'Issuance Status' from I/S workbook
        df_src = pd.read_excel(BytesIO(is_bytes), sheet_name="Issuance Status")
    else:
        # 👉 use external Issuance file
        if issuance_file is None:
            raise HTTPException(status_code=400, detail="issuance_file is required in external mode")
        df_src = pd.read_excel(issuance_file.file)

    # your existing function
    df_out = build_issuance_status(df_src)

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_out.to_excel(writer, sheet_name="Issuance Status", index=False)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="issuance_status.xlsx"'},
    )

from fastapi import Body

@app.post("/m1/device-id/export-json-xlsx")
async def export_device_id_json_xlsx(payload: dict = Body(...)):
    rows = payload.get("rows", [])
    df = pd.DataFrame(rows)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Redemption+DeviceID", index=False)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="redemption_with_device_ids_edited.xlsx"'
        },
    )



# ============================================================
# 1) PDF → Redemption Status (your existing logic)
# ============================================================
@app.post("/process-pdfs")
async def process_pdfs(files: List[UploadFile] = File(...)):
    """
    Accept multiple certificate PDFs, parse them, and return
    the final 'Redemption Status'-style table as JSON.
    """
    if not files:
        raise HTTPException(status_code=400, detail="No files uploaded")

    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            paths = []
            for f in files:
                path = os.path.join(tmpdir, f.filename)
                with open(path, "wb") as out:
                    out.write(await f.read())
                paths.append(path)

            # Use parse_one_pdf for each uploaded file
            # ✅ Sort PDFs like Colab (by date in filename + optional trailing sequence)
            paths = sorted(paths, key=filename_sort_key)
            
            # ✅ Parse in that order
            all_devices, all_certs = [], []
            for pdf in paths:
                d, c = parse_one_pdf(pdf)
                base_index_offset = len(all_devices)
                for row in c:
                    row["Device Index"] = base_index_offset + row["Device Index"]
                all_devices.extend(d)
                all_certs.extend(c)


        # Build DataFrames (similar to parse_pdf_folder)
        df_devices = pd.DataFrame(
            all_devices,
            columns=LABELS_FLAT + ["Source PDF", "Client Name"],
        )
        cert_cols = [
            "Device Index",
            "From Certificate ID",
            "To Certificate ID",
            "Number of Certificates",
            "Offset Attributes",
            "Period of Production",
            "Issuer",
            "Source PDF",
            "Client Name",
        ]
        df_certs = pd.DataFrame(all_certs, columns=cert_cols)

        df_devices_for_merge = (
            df_devices[LABELS_FLAT + ["Source PDF", "Client Name"]]
            .rename(columns={"Source PDF": "Device Source PDF"})
        )

        df_merged = pd.merge(
            df_certs,
            df_devices_for_merge,
            left_on="Device Index",
            right_index=True,
            how="left",
            suffixes=("_cert", "_dev"),
        ).drop(columns=["Device Index"])

        df_merged2 = df_merged.copy()
        df_merged2[["Start Date", "End Date"]] = df_merged2[
            "Period of Production"
        ].str.split(" - ", expand=True)

        df_sales = build_sales_dataframe(df_merged2)
        
        # ✅ Insert EXACTLY 2 blank rows between different PDFs (grouped by Source PDF)
        gap_rows = []
        prev_pdf = None
        
        for _, row in df_sales.iterrows():
            curr_pdf = str(row.get("Source PDF", ""))
            if prev_pdf is not None and curr_pdf != prev_pdf:
                gap_rows.append({c: "" for c in df_sales.columns})
                gap_rows.append({c: "" for c in df_sales.columns})
            gap_rows.append(row.to_dict())
            prev_pdf = curr_pdf
        
        # ✅ Re-number Sl cleanly (optional but recommended)
        sl = 1
        for r in gap_rows:
            is_blank = all(str(v).strip() == "" for v in r.values())
            if is_blank:
                r["Sl"] = ""
            else:
                r["Sl"] = sl
                sl += 1
        
        return JSONResponse(content=gap_rows)


    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))



# ============================================================
# 1b) PDFs + I/S master → updated I/S.xlsx
# ============================================================
from openpyxl import load_workbook

@app.post("/process-pdfs/is-xlsx")
async def process_pdfs_is_xlsx(
    files: List[UploadFile] = File(...),
    is_file: UploadFile = File(...)
):
    """
    Take certificate PDFs + an existing I/S workbook.
    Fill the 'Redemption Status' sheet and return XLSX.
    """
    import tempfile, os
    from io import BytesIO

    # --- Save and parse PDFs (same logic as /process-pdfs)
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            paths = []
            for f in files:
                path = os.path.join(tmpdir, f.filename)
                with open(path, "wb") as out:
                    out.write(await f.read())
                paths.append(path)

            all_devices, all_certs = [], []
            for pdf in paths:
                d, c = parse_one_pdf(pdf)
                base_index = len(all_devices)
                for row in c:
                    row["Device Index"] = base_index + row["Device Index"]
                all_devices.extend(d)
                all_certs.extend(c)

        df_devices = pd.DataFrame(all_devices, columns=LABELS_FLAT + ["Source PDF", "Client Name"])
        cert_cols = [
            "Device Index","From Certificate ID","To Certificate ID",
            "Number of Certificates","Offset Attributes","Period of Production",
            "Issuer","Source PDF","Client Name",
        ]
        df_certs = pd.DataFrame(all_certs, columns=cert_cols)

        df_devices_for_merge = df_devices[LABELS_FLAT + ["Source PDF", "Client Name"]]
        df_devices_for_merge = df_devices_for_merge.rename(columns={"Source PDF": "Device Source PDF"})

        df_merged = pd.merge(
            df_certs, df_devices_for_merge,
            left_on="Device Index", right_index=True,
            how="left"
        ).drop(columns=["Device Index"])

        df_merged2 = df_merged.copy()
        df_merged2[["Start Date","End Date"]] = df_merged2["Period of Production"].str.split(" - ", expand=True)

        df_sales = build_sales_dataframe(df_merged2)

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF parsing failed: {e}")

    # --- Load I/S workbook and overwrite Redemption Status sheet
    try:
        workbook_bytes = await is_file.read()
        wb = load_workbook(BytesIO(workbook_bytes))

        sheet_name = "Redemption Status"
        if sheet_name not in wb.sheetnames:
            raise HTTPException(status_code=400, detail="I/S workbook missing Redemption Status sheet")

        ws = wb[sheet_name]

        # Clear rows but keep headers
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

        # Write results starting row 2
        for r_idx, row in enumerate(df_sales.itertuples(index=False), start=2):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        out = BytesIO()
        wb.save(out)
        out.seek(0)

        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="IS_with_redemption.xlsx"'}
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to update I/S workbook: {e}")



# ============================================================
# 1.1) M-1 Parse Check
# ============================================================

@app.post("/m1/parse-check/json")
async def parse_check_json(payload: dict):
    rows = payload.get("rows", [])
    if not isinstance(rows, list) or not rows:
        raise HTTPException(status_code=400, detail="Missing rows")

    df = pd.DataFrame(rows)
    df_out = build_parse_check(df)

    # return JSON-safe
    return {"rows": df_out.replace({np.nan: ""}).to_dict(orient="records")}


# ============================================================
# 2) M-1 DeviceID: attach Device IDs from Excel files
# ============================================================

from fastapi import Body

@app.post("/m1/device-id/json")
async def process_device_id_json(payload: dict = Body(...)):
    """
    Return Redemption+DeviceID as JSON.
    Expects:
      {
        "sales_rows": [...],
        "registry_rows": [...]
      }
    """
    try:
        sales_rows = payload.get("sales_rows", [])
        registry_rows = payload.get("registry_rows", [])

        if not isinstance(sales_rows, list) or not sales_rows:
            raise HTTPException(status_code=400, detail="Missing or empty sales_rows")
        if not isinstance(registry_rows, list) or not registry_rows:
            raise HTTPException(status_code=400, detail="Missing or empty registry_rows")

        df_sales = pd.DataFrame(sales_rows)
        df_devreg = pd.DataFrame(registry_rows)

        print("=== /m1/device-id/json JSON DEBUG START ===")
        print("SALES COLUMNS:", df_sales.columns.tolist())
        print("REGISTRY COLUMNS:", df_devreg.columns.tolist())
        print("SALES SHAPE:", df_sales.shape)
        print("REGISTRY SHAPE:", df_devreg.shape)
        print("SALES HEAD:", df_sales.head(3).to_dict(orient="records"))
        print("REGISTRY HEAD:", df_devreg.head(3).to_dict(orient="records"))
        print("=== /m1/device-id/json JSON DEBUG END ===")

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not build DataFrames from JSON: {e}")

    try:
        import numpy as np

        tmp = df_sales.copy()
        tmp = tmp.replace(r"^\s*$", np.nan, regex=True)
        gap_mask = tmp.isna().all(axis=1)

        df_sales_nonblank = df_sales.loc[~gap_mask].copy()

        df_out_nonblank, fuzzy_log, unmatched = attach_device_ids(df_sales_nonblank, df_devreg)

        blank_row = {c: "" for c in df_out_nonblank.columns}
        out_rows = []
        it = iter(df_out_nonblank.to_dict(orient="records"))

        for is_gap in gap_mask.tolist():
            if is_gap:
                out_rows.append(blank_row.copy())
            else:
                out_rows.append(next(it))

        return JSONResponse(content=jsonable_encoder({
            "rows": out_rows,
            "fuzzy_log": [
                {"Sales Name": a, "Registry Name": b} for a, b in fuzzy_log
            ],
            "unmatched": unmatched.to_dict(orient="records"),
        }))

    except (KeyError, ValueError) as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Matching error: {e}")

@app.post("/m1/device-id/xlsx")
async def process_device_id_xlsx(
    sales_file: UploadFile = File(..., description="Redemption Status Excel"),
    registry_file: UploadFile = File(..., description="Device Registration Excel"),
):
    """
    Return Redemption+DeviceID as a downloadable Excel file.
    """
    try:
        df_sales = pd.read_excel(sales_file.file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read sales_file: {e}")

    try:
        df_devreg = pd.read_excel(registry_file.file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read registry_file: {e}")

    try:
        df_out, fuzzy_log, unmatched = attach_device_ids(df_sales, df_devreg)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Matching error: {e}")

    # Write to in-memory Excel
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_out.to_excel(writer, sheet_name="Redemption+DeviceID", index=False)
        if len(fuzzy_log) > 0:
            pd.DataFrame(fuzzy_log, columns=["Sales Name", "Registry Name"]).to_excel(
                writer, sheet_name="FuzzyMatches", index=False
            )
        if not unmatched.empty:
            unmatched.to_excel(writer, sheet_name="Unmatched", index=False)

    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": 'attachment; filename="redemption_with_device_ids.xlsx"'
        },
    )




from fastapi import UploadFile, File
from fastapi.responses import JSONResponse
import pandas as pd

from fastapi import UploadFile, File, HTTPException
from fastapi.responses import JSONResponse
from io import BytesIO
import pandas as pd
import traceback

@app.post("/m1/device-status/json")
async def device_status_json(
    devices_file: UploadFile = File(...),
    status_file: UploadFile = File(...),
):
    try:
        devices_bytes = await devices_file.read()
        status_bytes  = await status_file.read()

        df_devices = pd.read_excel(BytesIO(devices_bytes))
        df_status  = pd.read_excel(BytesIO(status_bytes))

        out = build_device_status_from_existing_status(df_devices, df_status, selling_price=4.5)
        return JSONResponse({"rows": out.to_dict(orient="records")})

    except KeyError as e:
        # column missing / pick() failed
        raise HTTPException(status_code=422, detail=f"Column missing: {str(e)}")
    except Exception as e:
        # keep this during debugging; later replace with proper logging
        tb = traceback.format_exc()
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}\n{tb}")




@app.post("/m1/device-status/xlsx")
async def device_status_xlsx(
    devices_file: UploadFile = File(..., description="Device Registration Excel"),
    vendor_file: UploadFile = File(..., description="Vendor Payment Excel"),
):
    try:
        df_devices = pd.read_excel(devices_file.file)
        df_vendor = pd.read_excel(vendor_file.file)
        out = build_device_status(df_devices, df_vendor, selling_price=4.5)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"DeviceStatus error: {e}")

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        out.to_excel(writer, sheet_name="DeviceStatus", index=False)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="device_status.xlsx"'},
    )



@app.post("/m1/issuance-status/json")
async def issuance_status_json(
    issuance_file: UploadFile = File(..., description="Device Issuance Status Excel"),
):
    try:
        df_src = pd.read_excel(issuance_file.file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read issuance_file: {e}")

    try:
        out = build_issuance_status(df_src)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"IssuanceStatus error: {e}")

    return JSONResponse({"rows": out.to_dict(orient="records")})


@app.post("/m1/issuance-status/xlsx")
async def issuance_status_xlsx(
    issuance_file: UploadFile = File(..., description="Device Issuance Status Excel"),
):
    try:
        df_src = pd.read_excel(issuance_file.file)
        out = build_issuance_status(df_src)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"IssuanceStatus error: {e}")

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        out.to_excel(writer, sheet_name="IssuanceStatus", index=False)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="issuance_status.xlsx"'},
    )



import pandas as pd
import numpy as np
from datetime import date, datetime

def df_records_json_safe(df: pd.DataFrame):
    df2 = df.copy()

    # Convert datetime-like columns to ISO strings
    for col in df2.columns:
        if pd.api.types.is_datetime64_any_dtype(df2[col]):
            df2[col] = df2[col].dt.strftime("%Y-%m-%d")

    def conv(x):
        if isinstance(x, (pd.Timestamp, datetime, date)):
            try:
                return x.isoformat()
            except Exception:
                return str(x)
        if x is None:
            return ""
        try:
            if pd.isna(x):
                return ""
        except Exception:
            pass
        return x

    # apply cell-wise without DataFrame.applymap
    df2 = df2.apply(lambda col: col.map(conv))
    return df2.to_dict(orient="records")


from io import BytesIO
import pandas as pd
import traceback
from fastapi import HTTPException, UploadFile, File
from fastapi.responses import JSONResponse

@app.post("/m1/cost-redemption/json")
async def cost_redemption_json(
    registration_file: UploadFile = File(...),
    redemption_file: UploadFile = File(...),
):
    try:
        reg_bytes = await registration_file.read()
        red_bytes = await redemption_file.read()

        registration_sheets = pd.read_excel(BytesIO(reg_bytes), sheet_name=None)
        df_red = pd.read_excel(BytesIO(red_bytes))

        df_out, stats = attach_redemption_cost(
            registration_sheets=registration_sheets,
            df_redemption=df_red,
            quarterly_start_year=25,
        )

        return JSONResponse({
            "rows": df_records_json_safe(df_out),
            "stats": stats,
        })

    except KeyError as e:
        raise HTTPException(status_code=422, detail=f"Missing column: {str(e)}")
    except Exception as e:
        tb = traceback.format_exc()
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}\n{tb}")


@app.post("/m1/cost-redemption/xlsx")
async def cost_redemption_xlsx(
    registration_file: UploadFile = File(..., description="Registration workbook with GRIT and Delnotic"),
    redemption_file: UploadFile = File(..., description="Redemption Status Excel"),
):
    """
    Return Redemption Status Excel with 'Cost of MWh' filled
    using Registration Data workbook (GRIT priority, Delnotic fallback).
    """
    try:
        reg_bytes = await registration_file.read()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read registration_file: {e}")

    try:
        df_red = pd.read_excel(redemption_file.file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read redemption_file: {e}")

    try:
        registration_sheets = pd.read_excel(BytesIO(reg_bytes), sheet_name=None)

        df_out, stats = attach_redemption_cost(
            registration_sheets=registration_sheets,
            df_redemption=df_red,
            quarterly_start_year=25,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"CostRedemption error: {e}")

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_out.to_excel(writer, sheet_name="Redemption Status", index=False)
        pd.DataFrame([stats]).to_excel(writer, sheet_name="Stats", index=False)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            'attachment; filename="redemption_with_cost_of_mwh.xlsx"'
        },
    )




from io import BytesIO
import pandas as pd
import traceback
from fastapi import UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse

def read_transfer_sheet_autheader(xlsx_bytes: bytes) -> pd.DataFrame:
    # preview without headers
    preview = pd.read_excel(
        BytesIO(xlsx_bytes),
        sheet_name="Transfer Status",
        header=None,
        nrows=20,
        engine="openpyxl",
    )

    header_row = None
    for i in range(len(preview)):
        row_vals = preview.iloc[i].astype(str).str.replace("\u00A0", " ").str.strip().str.lower().tolist()
        if "period starts" in row_vals:
            header_row = i
            break

    if header_row is None:
        raise ValueError(
            "Could not locate header row containing 'Period Starts'. "
            f"Preview rows: {preview.head(10).values.tolist()}"
        )

    # read full sheet using detected header row
    df = pd.read_excel(
        BytesIO(xlsx_bytes),
        sheet_name="Transfer Status",
        header=header_row,
        engine="openpyxl",
    )
    return df


@app.post("/m1/transfer-status/xlsx")
async def transfer_status_vintage(file: UploadFile = File(...)):
    try:
        b = await file.read()

        # ✅ header row 3 (0-index => 2)
        df_transfer = pd.read_excel(
            BytesIO(b),
            sheet_name="Transfer Status",
            header=2,
            engine="openpyxl",
        )
        df_transfer = read_transfer_sheet_autheader(b)

        df_out = attach_transfer_vintage(df_transfer, add_flags=False)

        out_buf = BytesIO()
        with pd.ExcelWriter(out_buf, engine="openpyxl") as writer:
            df_out.to_excel(writer, sheet_name="Transfer Status", index=False)
        out_buf.seek(0)

        return StreamingResponse(
            out_buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="transfer_status_with_vintage.xlsx"'},
        )

    except Exception as e:
        tb = traceback.format_exc()
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}\n{tb}")




from fastapi import UploadFile, File, HTTPException
from fastapi.responses import JSONResponse
from io import BytesIO
import pandas as pd
import traceback

def _q(dt: pd.Timestamp) -> int:
    return (dt.month - 1) // 3 + 1

def get_vintage(start_date, end_date) -> str:
    s = pd.to_datetime(start_date, errors="coerce")
    e = pd.to_datetime(end_date, errors="coerce")
    if pd.isna(s) or pd.isna(e):
        return ""
    return f"V{(s.year % 100):02d}Q{_q(s)}"

def highlight(start_date, end_date) -> bool:
    s = pd.to_datetime(start_date, errors="coerce")
    e = pd.to_datetime(end_date, errors="coerce")
    if pd.isna(s) or pd.isna(e):
        return False
    return _q(s) != _q(e)

@app.post("/m1/transfer-vintage/json")
async def transfer_vintage_json(file: UploadFile = File(...)):
    try:
        b = await file.read()

        # Auto-detect header row containing "Period Starts"
        preview = pd.read_excel(BytesIO(b), sheet_name="Transfer Status", header=None, nrows=20, engine="openpyxl")
        header_row = None
        for i in range(len(preview)):
            row = preview.iloc[i].astype(str).str.replace("\u00A0", " ").str.strip().str.lower().tolist()
            if "period starts" in row and "period ends" in row and "vintage" in row:
                header_row = i
                break
        if header_row is None:
            raise ValueError("Cannot find header row containing Period Starts/Ends/Vintage")

        df = pd.read_excel(BytesIO(b), sheet_name="Transfer Status", header=header_row, engine="openpyxl")
        df.columns = [str(c).replace("\u00A0", " ").strip() for c in df.columns]

        # Find IN/OUT columns by "base name" (handles 'Period Starts .1' etc.)
        import re
        def base(h):
            s = str(h).replace("\u00A0", " ").strip().lower()
            s = re.sub(r"\s*\.\d+\s*$", "", s)  # remove .1 /  .1
            return s

        def nth(name, k):
            m = [c for c in df.columns if base(c) == name.lower()]
            if len(m) < k:
                raise KeyError(f"Need {k} occurrence(s) of '{name}', found {len(m)}. Columns={list(df.columns)}")
            return m[k-1]

        in_start, in_end = nth("Period Starts", 1), nth("Period Ends", 1)
        out_start, out_end = nth("Period Starts", 2), nth("Period Ends", 2)

        in_v = [get_vintage(a, b) for a, b in zip(df[in_start], df[in_end])]
        out_v = [get_vintage(a, b) for a, b in zip(df[out_start], df[out_end])]

        in_h = [highlight(a, b) for a, b in zip(df[in_start], df[in_end])]
        out_h = [highlight(a, b) for a, b in zip(df[out_start], df[out_end])]

        return JSONResponse({"in_vintage": in_v, "out_vintage": out_v, "in_highlight": in_h, "out_highlight": out_h})

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}\n{traceback.format_exc()}")






@app.post("/m1/device-wise-sales/xlsx")
async def device_wise_sales_full_xlsx(
    issuance_file: UploadFile = File(..., description="Issuance Status Excel"),
    device_file: UploadFile = File(...,   description="Device Wise Sales Status Excel"),
    redemption_file: UploadFile = File(..., description="Redemption Status Excel"),
    transfer_file: UploadFile = File(...,   description="Transfer Status Excel"),
):
    try:
        issuance_df   = pd.read_excel(issuance_file.file)
        device_df     = pd.read_excel(device_file.file)
        redemption_df = pd.read_excel(redemption_file.file)
        tbytes = await transfer_file.read()
        transfer_df = read_transfer_sheet_autheader(tbytes)


        df_out = fill_device_wise_sales(
            device_df,
            issuance_df,
            redemption_df=redemption_df,
            transfer_df=transfer_df,
        )
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_out.to_excel(writer, sheet_name="Device Wise Sales Status", index=False)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="device_wise_sales_with_issued_sold.xlsx"'
        },
    )



@app.post("/m1/cogs/xlsx")
async def process_cogs_xlsx(
    cogs_file: UploadFile = File(..., description="COGS Excel"),
    sales_file: UploadFile = File(..., description="Device Wise Sales Status Excel"),
):
    """
    Upload:
      - COGS Excel
      - Device Wise Sales Status Excel

    Get back: COGS Excel with Registration Cost, Issued/Sold IREC,
              Issuance Cost, Redemption Cost filled.
    """
    try:
        cogs_df = pd.read_excel(cogs_file.file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read cogs_file: {e}")

    try:
        sales_df = pd.read_excel(sales_file.file)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read sales_file: {e}")

    try:
        merged = compute_cogs(cogs_df, sales_df)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"COGS computation error: {e}")

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        merged.to_excel(writer, sheet_name="COGS", index=False)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="COGS_with_costs.xlsx"'},
    )


from fastapi import Body
from .llm import ask_about_snapshot

@app.post("/m1/finance-chat/ask")
async def finance_chat_ask(payload: dict = Body(...)):
    question = str(payload.get("question", "")).strip()
    snapshot = payload.get("snapshot", {})
    previous_response_id = payload.get("previous_response_id")

    if not question:
        raise HTTPException(status_code=400, detail="Missing question")
    if not snapshot or "tabs" not in snapshot:
        raise HTTPException(status_code=400, detail="Missing snapshot tabs")

    try:
        out = ask_about_snapshot(
            question=question,
            snapshot=snapshot,
            previous_response_id=previous_response_id,
        )
        return out
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Finance chat failed: {e}")
